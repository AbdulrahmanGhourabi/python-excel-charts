from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference

wb = Workbook()
ws = wb.active


data = [
    ['Product', 'Sales', 'Profit', 'Expenses'],
    ['Apple', 100, 50, 30],
    ['Banana', 80, 40, 20],
    ['Cherry', 60, 30, 25],
    ['Date', 90, 45, 35],
]

for row in data:
    ws.append(row)


bar_chart = BarChart()
data_ref = Reference(ws, min_col=2, max_col=2, min_row=1, max_row=5)
cats = Reference(ws, min_col=1, min_row=2, max_row=5)
bar_chart.add_data(data_ref, titles_from_data=True)
bar_chart.set_categories(cats)
bar_chart.title = "Sales Bar Chart"
ws.add_chart(bar_chart, "E2")


line_chart = LineChart()
data_ref = Reference(ws, min_col=3, max_col=3, min_row=1, max_row=5)
line_chart.add_data(data_ref, titles_from_data=True)
line_chart.set_categories(cats)
line_chart.title = "Profit Line Chart"
ws.add_chart(line_chart, "E15")


pie_chart = PieChart()
data_ref = Reference(ws, min_col=4, max_col=4, min_row=2, max_row=5)
pie_chart.add_data(data_ref, titles_from_data=False)
pie_chart.set_categories(cats)
pie_chart.title = "Expenses Pie Chart"
ws.add_chart(pie_chart, "E28")

wb.save("charts.xlsx")
print("Excel file with 3 charts created!")